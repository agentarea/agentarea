"""Refs-only artifact validation executed in the canonical task sandbox."""

from __future__ import annotations

import json
from collections.abc import Awaitable, Callable, Sequence
from pathlib import PurePosixPath
from typing import Any

from agentarea_common.artifacts import (
    WorkspaceRepository,
    WorkspaceValidationError,
    normalize_workspace_path,
)

from ..models import (
    ArtifactValidationEvidence,
    ArtifactValidationIssue,
    ArtifactValidationRequest,
    ArtifactValidationResult,
    CapabilityUnavailableResult,
    RuntimeDiscoveryResult,
)
from .runtime_discovery import require_runtime_capability

SandboxExecutor = Callable[[str], Awaitable[dict[str, Any]]]

_VALIDATORS = {
    ".py": "py_compile",
    ".xlsx": "openpyxl",
    ".docx": "python-docx",
    ".pptx": "python-pptx",
    ".pdf": "pypdf",
    ".html": "playwright",
    ".htm": "playwright",
}
_PACKAGE_KEYS = {
    "openpyxl": ("openpyxl",),
    "python-docx": ("python-docx", "docx"),
    "python-pptx": ("python-pptx", "pptx"),
    "pypdf": ("pypdf",),
}


def _validator_for(path: str) -> str | None:
    return _VALIDATORS.get(PurePosixPath(path).suffix.casefold())


def _unavailable(
    *, generation: int, capability: str, runtime: RuntimeDiscoveryResult
) -> ArtifactValidationResult:
    manifest = runtime.manifest
    return ArtifactValidationResult(
        state="unavailable",
        generation=generation,
        capability_unavailable=CapabilityUnavailableResult(
            capability=capability,
            runtime_version=manifest.image_version if manifest else None,
        ),
        issues=[
            ArtifactValidationIssue(
                path="",
                validator=capability,
                code="capability_unavailable",
                message=f"Required validation capability is unavailable: {capability}",
            )
        ],
    )


def _build_validator_command(paths: Sequence[str], *, browser_available: bool) -> str:
    """Build a fixed validator program; interpolated paths are JSON literals only."""
    payload = json.dumps(list(paths), ensure_ascii=True)
    browser = "True" if browser_available else "False"
    return f"""python - <<'AGENTAREA_VALIDATOR'
import hashlib
import json
import os
import py_compile
import tempfile
from pathlib import Path

paths = {payload}
browser_available = {browser}
results = []

def safe_error(exc, validator_name):
    if isinstance(exc, py_compile.PyCompileError):
        value = getattr(exc, 'exc_value', None)
        line = getattr(value, 'lineno', None)
        message = getattr(value, 'msg', None) or 'Python compilation failed'
        return f'{{type(value).__name__ if value else "SyntaxError"}} at line {{line or "unknown"}}: {{message}}'
    if isinstance(exc, FileNotFoundError):
        return 'Declared artifact does not exist'
    return f'{{type(exc).__name__}}: artifact could not be opened by {{validator_name}}'

for relative_path in paths:
    path = Path(relative_path)
    validator = {{
        '.py': 'py_compile',
        '.xlsx': 'openpyxl',
        '.docx': 'python-docx',
        '.pptx': 'python-pptx',
        '.pdf': 'pypdf',
        '.html': 'playwright',
        '.htm': 'playwright',
    }}.get(path.suffix.casefold(), 'existence')
    try:
        if not path.is_file():
            raise FileNotFoundError('declared artifact does not exist')
        if validator == 'py_compile':
            handle, compiled_path = tempfile.mkstemp(prefix='agentarea-validation-', suffix='.pyc')
            os.close(handle)
            try:
                py_compile.compile(str(path), cfile=compiled_path, doraise=True)
            finally:
                Path(compiled_path).unlink(missing_ok=True)
        elif validator == 'openpyxl':
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            workbook.close()
        elif validator == 'python-docx':
            from docx import Document
            Document(path)
        elif validator == 'python-pptx':
            from pptx import Presentation
            Presentation(path)
        elif validator == 'pypdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            len(reader.pages)
        elif validator == 'playwright':
            if not browser_available:
                raise RuntimeError('browser capability is unavailable')
            from playwright.sync_api import sync_playwright
            with sync_playwright() as playwright:
                instance = playwright.chromium.launch(headless=True)
                try:
                    page = instance.new_page()
                    page.goto(path.resolve().as_uri(), wait_until='load')
                    page.content()
                finally:
                    instance.close()
        results.append({{'path': relative_path, 'validator': validator, 'status': 'passed'}})
    except ModuleNotFoundError as exc:
        results.append({{
            'path': relative_path,
            'validator': validator,
            'status': 'unavailable',
            'code': 'capability_unavailable',
            'message': f'missing runtime module: {{exc.name}}',
        }})
    except Exception as exc:
        results.append({{
            'path': relative_path,
            'validator': validator,
            'status': 'failed',
            'code': 'validation_failed',
            'message': safe_error(exc, validator)[:1000],
        }})

print(json.dumps({{'results': results}}, ensure_ascii=True, separators=(',', ':')))
AGENTAREA_VALIDATOR"""


async def validate_workspace_artifacts(
    request: ArtifactValidationRequest,
    *,
    repository: WorkspaceRepository,
    runtime: RuntimeDiscoveryResult,
    execute_in_sandbox: SandboxExecutor,
) -> ArtifactValidationResult:
    """Validate artifact identities without moving their bodies through Temporal or Redis."""
    manifest_ref = await repository.current_manifest_ref(request.workspace_id, request.task_id)
    objects = await repository.list(request.workspace_id, request.task_id)
    generation = max(
        (item.generation for item in objects),
        default=manifest_ref.generation,
    )
    by_path = {item.path: item for item in objects}

    declared: list[str] = []
    for raw_path in request.declared_paths:
        try:
            path = normalize_workspace_path(raw_path)
        except WorkspaceValidationError as exc:
            return ArtifactValidationResult(
                state="failed",
                generation=generation,
                issues=[
                    ArtifactValidationIssue(
                        path=str(raw_path),
                        validator="existence",
                        code="invalid_artifact_path",
                        message=str(exc),
                    )
                ],
            )
        if path not in declared:
            declared.append(path)

    missing = [path for path in declared if path not in by_path]
    if missing:
        return ArtifactValidationResult(
            state="failed",
            generation=generation,
            issues=[
                ArtifactValidationIssue(
                    path=path,
                    validator="existence",
                    code="artifact_missing",
                    message="Declared artifact does not exist in the committed task workspace",
                )
                for path in missing
            ],
        )

    selected = list(declared)
    selected.extend(
        item.path
        for item in objects
        if item.path not in selected
        and not item.path.startswith(("inputs/", "skills/", ".agentarea/"))
        and _validator_for(item.path) is not None
    )
    if not selected:
        return ArtifactValidationResult(state="no_artifacts", generation=generation)

    validators = {path: _validator_for(path) or "existence" for path in selected}
    manifest = runtime.manifest
    if manifest is None:
        return _unavailable(
            generation=generation,
            capability="runtime_manifest",
            runtime=runtime,
        )

    if any(validator == "playwright" for validator in validators.values()):
        unavailable = require_runtime_capability(runtime, "browser")
        if unavailable is not None:
            return ArtifactValidationResult(
                state="unavailable",
                generation=generation,
                capability_unavailable=unavailable,
                issues=[
                    ArtifactValidationIssue(
                        path=path,
                        validator="playwright",
                        code="capability_unavailable",
                        message="HTML smoke validation requires the browser capability",
                    )
                    for path, validator in validators.items()
                    if validator == "playwright"
                ],
            )

    packages = {name.casefold() for name in manifest.packages}
    for validator in sorted(set(validators.values())):
        expected = _PACKAGE_KEYS.get(validator)
        if expected and not any(name in packages for name in expected):
            return _unavailable(
                generation=generation,
                capability=f"python_package:{validator}",
                runtime=runtime,
            )

    outcome = await execute_in_sandbox(
        _build_validator_command(
            selected,
            browser_available=require_runtime_capability(runtime, "browser") is None,
        )
    )
    if not outcome.get("success"):
        raise RuntimeError(str(outcome.get("error") or outcome.get("result") or "validator failed"))

    # The sandbox checkout lease pins the hydrated snapshot while the validator
    # runs. Re-read identities after writeback to close the small window between
    # initial discovery and checkout; command/output bookkeeping may advance the
    # generation, but the artifacts that were actually validated must not change.
    final_objects = {
        item.path: item for item in await repository.list(request.workspace_id, request.task_id)
    }
    for path in selected:
        before = by_path[path]
        after = final_objects.get(path)
        if after is None or (after.sha256, after.size, after.object_uri) != (
            before.sha256,
            before.size,
            before.object_uri,
        ):
            raise RuntimeError(f"artifact changed during validation: {path}")
    try:
        raw = json.loads(str(outcome.get("result") or ""))
        results = raw["results"]
        if not isinstance(results, list):
            raise TypeError("results must be a list")
    except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise RuntimeError("sandbox validator returned malformed output") from exc

    issues: list[ArtifactValidationIssue] = []
    evidence: list[ArtifactValidationEvidence] = []
    saw_unavailable = False
    returned_paths: set[str] = set()
    for value in results:
        if not isinstance(value, dict):
            raise RuntimeError("sandbox validator returned a malformed result item")
        path = str(value.get("path") or "")
        validator = str(value.get("validator") or "unknown")
        if path in returned_paths or path not in validators or validator != validators[path]:
            raise RuntimeError("sandbox validator returned mismatched artifact identities")
        returned_paths.add(path)
        status = value.get("status")
        item = by_path.get(path)
        if status == "passed" and item is not None:
            evidence.append(
                ArtifactValidationEvidence(
                    path=path,
                    validator=validator,
                    sha256=item.sha256,
                    size=item.size,
                )
            )
            continue
        if status == "unavailable":
            saw_unavailable = True
        issues.append(
            ArtifactValidationIssue(
                path=path,
                validator=validator,
                code=str(value.get("code") or "validation_failed"),
                message=str(value.get("message") or "Artifact validation failed")[:1000],
            )
        )

    if returned_paths != set(selected):
        raise RuntimeError("sandbox validator omitted artifact results")

    if saw_unavailable:
        return ArtifactValidationResult(
            state="unavailable",
            generation=generation,
            evidence=evidence,
            issues=issues,
            capability_unavailable=CapabilityUnavailableResult(capability="artifact_validator"),
        )
    if issues:
        return ArtifactValidationResult(
            state="failed", generation=generation, evidence=evidence, issues=issues
        )
    return ArtifactValidationResult(
        state="passed", generation=generation, evidence=evidence
    )
